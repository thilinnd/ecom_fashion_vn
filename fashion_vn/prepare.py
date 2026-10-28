"""Audit XLSX/images, group duplicates, export weak labels and blind review sheets."""
import argparse
import collections
import csv
import hashlib
import json
from pathlib import Path
import random

import openpyxl
from PIL import Image, ImageOps
from .text import extract, category_hint, LEXICON

def write_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')

def read_rows(path):
    return [json.loads(line) for line in Path(path).read_text(encoding='utf-8').splitlines() if line.strip()]

def write_csv(path, rows, fields):
    with Path(path).open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)

def prepare(args):
    source, root, out = Path(args.xlsx).resolve(), Path(args.image_root).resolve(), Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    sheet = wb[args.sheet] if args.sheet else wb.active
    it = sheet.iter_rows(values_only=True)
    headers = list(next(it))
    required = ['product_id', 'product', 'img_local_path']
    if any(c not in headers for c in required):
        raise ValueError(f'Required columns: {required}; found {headers}')
    positions = [headers.index(c) for c in required]
    records, ids = [], set()
    for row in it:
        pid, title, rel = [row[i] for i in positions]
        if pid is None and title is None:
            continue
        pid = str(pid)
        if pid in ids:
            raise ValueError(f'Duplicate product_id: {pid}')
        ids.add(pid)
        norm, tokens, attrs, evidence, flags = extract(title)
        path = (root / str(rel or '')).resolve()
        status, sha, phash, size = 'missing', '', '', None
        if not path.is_relative_to(root):
            status = 'outside_image_root'
        elif path.is_file():
            try:
                raw = path.read_bytes()
                sha = hashlib.sha256(raw).hexdigest()
                with Image.open(path) as im:
                    im = ImageOps.exif_transpose(im).convert('RGB')
                    im.load()
                    size = list(im.size)
                    small = im.resize((9, 8)).convert('L')
                    p = list(small.getdata())
                    bits = [p[y*9+x] > p[y*9+x+1] for y in range(8) for x in range(8)]
                    phash = f'{sum(int(v) << i for i, v in enumerate(bits)):016x}'
                status = 'ok'
            except Exception as exc:
                status = 'invalid:' + type(exc).__name__
        if not norm:
            flags.append('empty_title')
        category = category_hint(norm)
        if category not in {'dress', 'top', 'pants', 'skirt', 'jacket'}:
            flags.append('outside_original_five_categories')
        records.append(dict(product_id=pid, title_raw=str(title or ''), title_normalized=norm,
                            image_path=str(path), image_status=status, image_sha256=sha,
                            image_dhash=phash, image_size=size, category=category,
                            category_source='title_rule_unverified', tokens=tokens,
                            title_attributes=attrs, title_evidence=evidence, flags=flags))
        if len(records) % 2000 == 0:
            print(f'Audited {len(records)} products', flush=True)
    wb.close()
    # Connected components prevent transitive title/image duplicates leaking.
    parent = list(range(len(records)))
    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i
    seen = {}
    for i, r in enumerate(records):
        keys = [('title', r['title_normalized']), ('sha', r['image_sha256']),
                ('dhash_exact', r['image_dhash'])]
        for typ, val in keys:
            if not val:
                continue
            if (typ, val) in seen:
                parent[find(i)] = find(seen[typ, val])
            seen[typ, val] = i
    groups = collections.defaultdict(list)
    for i, r in enumerate(records):
        groups[find(i)].append(r)
    for group in groups.values():
        gid = hashlib.sha256('|'.join(sorted(r['product_id'] for r in group)).encode()).hexdigest()
        bucket = int(hashlib.sha256((str(args.seed) + gid).encode()).hexdigest()[:8], 16) / 2**32
        split = 'train' if bucket < .8 else 'val' if bucket < .9 else 'test'
        for r in group:
            r.update(group_id=gid, split=split)
    with (out / 'manifest.jsonl').open('w', encoding='utf-8') as f:
        for r in records:
            f.write(json.dumps(r, ensure_ascii=False) + '\n')
    df = collections.Counter(t for r in records if r['split'] == 'train' for t in set(r['tokens']))
    candidate_rows = [{'token': t, 'train_document_frequency': n,
                       'include': int(n >= args.min_df), 'review_note': ''}
                      for t, n in df.most_common()]
    write_csv(out / 'vocabulary_review.csv', candidate_rows,
              ['token', 'train_document_frequency', 'include', 'review_note'])
    audit = {
        'source': str(source), 'n_products': len(records), 'n_duplicate_groups': len(groups),
        'rows_beyond_unique_raw_titles': len(records) - len({r['title_raw'] for r in records}),
        'rows_beyond_unique_normalized_titles': len(records) - len({r['title_normalized'] for r in records}),
        'image_status': dict(collections.Counter(r['image_status'] for r in records)),
        'category_hints_unverified': dict(collections.Counter(r['category'] for r in records)),
        'splits': dict(collections.Counter(r['split'] for r in records)),
        'flags': dict(collections.Counter(f for r in records for f in r['flags'])),
        'products_with_title_attributes': sum(bool(r['title_attributes']) for r in records),
        'attribute_mention_counts': dict(collections.Counter(c for r in records for c in r['title_attributes'])),
        'train_vocabulary_size_min_df': sum(n >= args.min_df for n in df.values()),
        'seed': args.seed,
        'split_policy': '80/10/10 hash of connected duplicate groups: normalized title, SHA256, exact dHash',
        'limitations': ['Category and title attributes are unverified weak labels.',
                       'Exact dHash groups are conservative; near duplicates and seller leakage need additional review.',
                       'Unmentioned attributes are unknown, not negative.',
                       'No learned image predictions are produced by prepare.'],
    }
    write_json(out / 'audit.json', audit)
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = 'title_weak_labels'
    cols = ['product_id', 'title_raw', 'title_normalized', 'category', 'image_path', 'image_status',
            'split', 'flags', 'label_source', *LEXICON]
    ws.append(cols)
    for row_number, r in enumerate(records, start=2):
        values = [r[k] for k in cols[:7]] + [';'.join(r['flags']), 'title_rule_unverified']
        values += [';'.join(r['title_attributes'].get(k, [])) for k in LEXICON]
        ws.append(values)
        for column in range(1, len(cols)+1):
            cell = ws.cell(row=row_number, column=column)
            if isinstance(cell.value, str):
                cell.data_type = 's'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    book.save(out / 'title_weak_labels.xlsx')
    # Blind gold sheets deliberately omit title and machine suggestions.
    rng = random.Random(args.seed)
    samples = []
    for split in ('val', 'test'):
        buckets = collections.defaultdict(list)
        for r in records:
            if r['split'] == split and r['image_status'] == 'ok':
                buckets[r['category']].append(r)
        for pool in buckets.values():
            rng.shuffle(pool)
        chosen, seen_groups = [], set()
        while len(chosen) < args.review_size and any(buckets.values()):
            for cat in sorted(buckets):
                if buckets[cat] and len(chosen) < args.review_size:
                    r = buckets[cat].pop()
                    if r['group_id'] not in seen_groups:
                        chosen.append(r)
                        seen_groups.add(r['group_id'])
        samples.extend(chosen)
    review = [{'product_id': r['product_id'], 'split': r['split'], 'image_path': r['image_path'],
               'category_gold': '', 'image_quality': '', 'multiple_products': '',
               'color_gold': '', 'sleeve_gold': '', 'neckline_gold': '', 'pattern_gold': '',
               'fit_gold': '', 'length_gold': '', 'pant_cut_gold': '', 'waist_gold': '',
               'set_type_gold': '', 'notes': '', 'annotator': ''} for r in samples]
    fields = list(review[0]) if review else ['product_id', 'split', 'image_path', 'notes']
    write_csv(out / 'annotation_blind.csv', review, fields)
    write_csv(out / 'gold_long_template.csv', [], ['product_id', 'split', 'concept', 'attribute', 'label', 'annotator'])
    print(json.dumps(audit, ensure_ascii=True, indent=2))

def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--xlsx', default='data/fashion_raw.xlsx')
    p.add_argument('--image-root', default='data')
    p.add_argument('--sheet')
    p.add_argument('--out', default='outputs/prepared')
    p.add_argument('--min-df', type=int, default=5)
    p.add_argument('--seed', type=int, default=42)
    p.add_argument('--review-size', type=int, default=200, help='Products per val/test split')
    prepare(p.parse_args())

if __name__ == '__main__':
    main()
