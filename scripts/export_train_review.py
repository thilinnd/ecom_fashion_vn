"""Export the current VSE training population to editable review files.

The export mirrors the eligibility and per-category vocabulary filtering in
``fashion_vn.train``.  It does not modify the manifest or trained models.
"""

from __future__ import annotations

import argparse
import collections
import csv
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

MODEL_CATEGORIES = ["top", "dress", "pants", "skirt", "jacket", "bag", "hat", "set"]
EXCLUDED_FLAGS = {"listing_multiple_variants", "possible_negation_review"}


def read_rows(path: Path) -> list[dict]:
    with path.open(encoding="utf-8") as handle:
        return [json.loads(line) for line in handle if line.strip()]


def classify_train_rows(rows: list[dict], min_df: int) -> tuple[list[dict], list[dict]]:
    train_rows = [row for row in rows if row["split"] == "train"]
    vocabularies: dict[str, set[str]] = {}

    for category in MODEL_CATEGORIES:
        eligible = [
            row
            for row in train_rows
            if row["category"] == category
            and row["image_status"] == "ok"
            and not EXCLUDED_FLAGS.intersection(row["flags"])
        ]
        document_frequency = collections.Counter(
            token for row in eligible for token in set(row["tokens"])
        )
        vocabularies[category] = {
            token for token, count in document_frequency.items() if count >= min_df
        }

    included, excluded = [], []
    for row in train_rows:
        category = row["category"]
        reasons = []
        if category not in MODEL_CATEGORIES:
            reasons.append("category_not_modelled")
        if row["image_status"] != "ok":
            reasons.append("invalid_image")
        flagged = sorted(EXCLUDED_FLAGS.intersection(row["flags"]))
        if flagged:
            reasons.append("excluded_flag:" + ";".join(flagged))
        if not reasons and not vocabularies[category].intersection(row["tokens"]):
            reasons.append("no_token_meets_category_min_df")

        exported = {
            "product_id": row["product_id"],
            "include_in_retrain": 0 if reasons else 1,
            "current_status": "excluded" if reasons else "included",
            "exclusion_reason": ";".join(reasons),
            "category_original": category,
            "category_override": "",
            "title_raw": row["title_raw"],
            "title_override": "",
            "image_path": row["image_path"],
            "image_status": row["image_status"],
            "flags": ";".join(row["flags"]),
            "tokens": json.dumps(row["tokens"], ensure_ascii=False),
            "title_attributes": json.dumps(row["title_attributes"], ensure_ascii=False),
            "duplicate_group_id": row["group_id"],
            "review_note": "",
        }
        (excluded if reasons else included).append(exported)

    return included, excluded


def write_csv(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def append_sheet(workbook: openpyxl.Workbook, name: str, rows: list[dict]) -> None:
    sheet = workbook.create_sheet(name)
    columns = list(rows[0])
    sheet.append(columns)
    for row in rows:
        sheet.append([row[column] for column in columns])

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F81BD")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 13
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 15
    sheet.column_dimensions["D"].width = 34
    sheet.column_dimensions["E"].width = 18
    sheet.column_dimensions["F"].width = 18
    sheet.column_dimensions["G"].width = 75
    sheet.column_dimensions["H"].width = 50
    sheet.column_dimensions["I"].width = 55
    sheet.column_dimensions["O"].width = 45

    image_column = columns.index("image_path") + 1
    for row_number in range(2, sheet.max_row + 1):
        cell = sheet.cell(row_number, image_column)
        cell.hyperlink = cell.value
        cell.style = "Hyperlink"


def write_workbook(path: Path, included: list[dict], excluded: list[dict]) -> None:
    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "summary"
    summary.append(["Item", "Count"])
    summary.append(["Current model-training products", len(included)])
    summary.append(["Excluded train-split products", len(excluded)])
    summary.append(["All train-split products", len(included) + len(excluded)])
    summary.append(["Instruction", "Edit only include_in_retrain, category_override, title_override, and review_note."])
    summary.append(["Safety", "This workbook does not change manifest.jsonl until an explicit import step is run."])
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F81BD")
    summary.column_dimensions["A"].width = 36
    summary.column_dimensions["B"].width = 100

    append_sheet(workbook, "model_train", included)
    append_sheet(workbook, "excluded_train", excluded)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", type=Path, default=Path("outputs/prepared/manifest.jsonl"))
    parser.add_argument("--min-df", type=int, default=5)
    parser.add_argument(
        "--csv-output",
        type=Path,
        default=Path("outputs/prepared/train_model_current.csv"),
    )
    parser.add_argument(
        "--xlsx-output",
        type=Path,
        default=Path("outputs/prepared/train_review.xlsx"),
    )
    args = parser.parse_args()

    included, excluded = classify_train_rows(read_rows(args.manifest), args.min_df)
    write_csv(args.csv_output, included)
    write_workbook(args.xlsx_output, included, excluded)
    print(f"Current model train: {len(included):,}")
    print(f"Excluded train rows: {len(excluded):,}")
    print(f"CSV: {args.csv_output.resolve()}")
    print(f"Workbook: {args.xlsx_output.resolve()}")


if __name__ == "__main__":
    main()
