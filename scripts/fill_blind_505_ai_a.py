"""Fill blind_505_before_AI_A.xlsx from an image-only AI_A review."""
from pathlib import Path
import shutil

from openpyxl import load_workbook


COLUMNS = [
    "category_gold", "image_quality", "multiple_products", "color_gold",
    "sleeve_gold", "neckline_gold", "pattern_gold", "fit_gold", "length_gold",
    "pant_cut_gold", "waist_gold", "notes", "annotator",
]
NA = "not_applicable"
SPECIAL = {"unknown", "not_visible", NA, "other"}
ALLOWED = {
    "category_gold": {"top", "dress", "pants", "skirt", "jacket", "bag", "hat", "shoe", "set", "fabric", "other", "unknown"},
    "image_quality": {"good", "usable", "poor", "unusable"},
    "multiple_products": {"yes", "no", "unknown"},
    "color_gold": {"đen", "trắng", "đỏ", "hồng", "xanh_dương", "xanh_lá", "xanh_navy", "vàng", "tím", "nâu", "xám", "be", "kem"} | SPECIAL,
    "sleeve_gold": {"tay_dài", "tay_ngắn", "tay_lỡ", "không_tay", "tay_phồng", "tay_loe", "tay_raglan"} | SPECIAL,
    "neckline_gold": {"cổ_tròn", "cổ_v", "cổ_vuông", "cổ_lọ", "cổ_yếm", "cổ_đổ", "cổ_bẻ", "trễ_vai", "hai_dây", "một_vai"} | SPECIAL,
    "pattern_gold": {"trơn", "kẻ_sọc", "caro", "chấm_bi", "hoa", "in_hình", "loang", "da_báo"} | SPECIAL,
    "fit_gold": {"rộng", "ôm", "suông"} | SPECIAL,
    "length_gold": {"maxi", "midi", "mini", "croptop"} | SPECIAL,
    "pant_cut_gold": {"ống_rộng", "ống_suông", "ống_loe", "ống_đứng", "ống_bó"} | SPECIAL,
    "waist_gold": {"cạp_cao", "cạp_thấp"} | SPECIAL,
    "annotator": {"AI_A"},
}


def row(*values):
    if len(values) != len(COLUMNS):
        raise ValueError(f"Expected {len(COLUMNS)} values, got {len(values)}")
    return dict(zip(COLUMNS, values))


ANNOTATIONS = {
    11060: row("jacket", "usable", "yes", "nâu", "tay_dài", "other", "trơn", "rộng", "other", NA, NA, "Cùng một áo khoác có mũ trên hai người mẫu", "AI_A"),
    1092: row("pants", "good", "no", "xanh_navy", NA, NA, "trơn", "suông", NA, "ống_đứng", "cạp_cao", "Quần denim màu xanh đậm", "AI_A"),
    4173: row("set", "good", "yes", "vàng;đỏ", "không_tay", "cổ_tròn", "trơn", "ôm", "croptop", "ống_bó", "cạp_cao", "Bộ đồ thể thao màu vàng; ảnh phụ minh họa áo màu đỏ", "AI_A"),
    8380: row("skirt", "good", "yes", "đen", NA, NA, "trơn", "ôm", "mini", NA, "cạp_cao", "Chân váy denim được phối với áo", "AI_A"),
    9739: row("top", "good", "yes", "trắng", "không_tay", "cổ_đổ", "trơn", "suông", "other", NA, NA, "Ảnh phối áo hai dây với áo khoác và quần đen", "AI_A"),
    5258: row("bag", "good", "no", "xám", NA, NA, "kẻ_sọc", NA, NA, NA, NA, "Túi đeo hông có ba sọc dọc ở mặt trước", "AI_A"),
    5706: row("bag", "good", "yes", "đen;đỏ", NA, NA, "in_hình", NA, NA, NA, NA, "Bộ ba lô và túi nhỏ trẻ em có hình in", "AI_A"),
    3158: row("skirt", "good", "yes", "xanh_dương", NA, NA, "trơn", "suông", "midi", NA, "cạp_cao", "Cùng một chân váy denim ở mặt trước và sau", "AI_A"),
    5779: row("hat", "good", "yes", "unknown", NA, NA, "in_hình", NA, NA, NA, NA, "Nhiều biến thể màu của mũ lưỡi trai", "AI_A"),
    2456: row("jacket", "good", "yes", "xanh_dương", "tay_dài", "other", "trơn;caro", "rộng", "other", NA, NA, "Áo khoác denim có phần mũ và lớp phối caro", "AI_A"),
    12339: row("pants", "good", "no", "vàng", NA, NA, "trơn", "ôm", NA, NA, "cạp_thấp", "Quần ngắn ánh kim; loại ống quần không áp dụng", "AI_A"),
    3937: row("set", "good", "yes", "trắng;đen", "tay_ngắn", "cổ_tròn", "in_hình", "suông", "other", NA, "unknown", "Bộ áo thun và quần ngắn", "AI_A"),
    8626: row("skirt", "good", "no", "xanh_dương", NA, NA, "trơn", "ôm", "midi", NA, "cạp_cao", "Chân váy denim dáng bút chì", "AI_A"),
    12517: row("set", "usable", "yes", "đỏ;xanh_lá;xanh_dương", "tay_dài", "cổ_tròn", "trơn", "suông", "other", "ống_rộng", "not_visible", "Nhiều biến thể của bộ áo dài truyền thống và quần", "AI_A"),
    4100: row("set", "usable", "yes", "đỏ;đen", "không_tay", "cổ_tròn", "trơn", "ôm", "croptop", NA, "cạp_cao", "Bộ áo thể thao và quần ngắn ở nhiều góc nhìn", "AI_A"),
    5298: row("bag", "good", "yes", "xám", NA, NA, "trơn", NA, NA, NA, NA, "Túi đeo chéo được chụp trên người mẫu", "AI_A"),
    7808: row("dress", "good", "yes", "xanh_navy", "tay_ngắn", "cổ_tròn", "trơn", "suông", "maxi", NA, "other", "Váy dài được phối thêm túi xách", "AI_A"),
    5912: row("hat", "good", "yes", "unknown", NA, NA, "in_hình", NA, NA, NA, NA, "Nhiều biến thể màu của mũ len có chữ", "AI_A"),
    2086: row("jacket", "good", "yes", "xám", "tay_dài", "other", "trơn", "rộng", "other", NA, NA, "Cùng một áo hoodie khóa kéo ở mặt trước và sau", "AI_A"),
    4351: row("set", "good", "yes", "unknown", "tay_ngắn", "cổ_tròn", "trơn", "ôm", "croptop", "ống_bó", "cạp_cao", "Nhiều biến thể màu của bộ đồ tập", "AI_A"),
    12732: row("set", "good", "yes", "unknown", "tay_ngắn", "cổ_tròn", "caro", "suông", "other", "ống_suông", "unknown", "Nhiều biến thể màu của bộ áo và quần trẻ em", "AI_A"),
    8031: row("skirt", "good", "no", "nâu", NA, NA, "trơn", "suông", "mini", NA, "cạp_cao", "Chân váy xếp ly có quần bảo hộ bên trong", "AI_A"),
    9588: row("top", "good", "yes", "unknown", "tay_lỡ", "cổ_lọ", "trơn", "suông", "other", NA, NA, "Nhiều biến thể màu của áo tay phồng", "AI_A"),
    11208: row("jacket", "good", "no", "đỏ", "không_tay", "other", "trơn", "suông", "other", NA, NA, "Áo ghi-lê cổ đứng với khuy cài trang trí", "AI_A"),
    5262: row("bag", "good", "yes", "xám;hồng", NA, NA, "trơn", NA, NA, NA, NA, "Túi du lịch được trình bày cùng phụ kiện và biến thể màu", "AI_A"),
    7892: row("set", "usable", "yes", "hồng", "tay_dài", "cổ_v", "hoa", "suông", "maxi", NA, "cạp_cao", "Bộ áo và chân váy dài hoa ở nhiều khung hình", "AI_A"),
    6383: row("hat", "good", "yes", "unknown", NA, NA, "in_hình", NA, NA, NA, NA, "Nhiều biến thể mũ bucket trẻ em hình động vật", "AI_A"),
    4773: row("jacket", "good", "yes", "be", "tay_dài", "cổ_bẻ", "trơn", "suông", "other", NA, NA, "Blazer có các biến thể màu minh họa ở bên phải", "AI_A"),
    1219: row("pants", "good", "yes", "xanh_dương", NA, NA, "trơn", "suông", NA, NA, "cạp_cao", "Nhiều biến thể quần denim ngắn; loại ống quần không áp dụng", "AI_A"),
    4269: row("set", "good", "yes", "xanh_lá", "không_tay", "hai_dây", "trơn", "ôm", "croptop", "ống_bó", "cạp_cao", "Cùng một bộ đồ thể thao ở nhiều góc nhìn", "AI_A"),
    9021: row("skirt", "usable", "yes", "đỏ", NA, NA, "caro", "suông", "mini", NA, "cạp_cao", "Chân váy caro được phối với áo len", "AI_A"),
    10460: row("top", "good", "yes", "xanh_navy", "tay_ngắn", "cổ_tròn", "in_hình", "rộng", "other", NA, NA, "Áo thun dáng rộng được trình bày ở mặt trước và sau", "AI_A"),
    4154: row("set", "good", "yes", "xanh_navy", "tay_ngắn", "cổ_bẻ", "in_hình", "suông", "other", NA, "unknown", "Bộ áo polo và quần ngắn thể thao", "AI_A"),
    5077: row("bag", "good", "yes", "nâu;be", NA, NA, "trơn", NA, NA, NA, NA, "Nhiều túi xách với kiểu dáng và tông màu khác nhau", "AI_A"),
    7451: row("dress", "usable", "yes", "kem", "không_tay", "cổ_yếm", "hoa", "suông", "maxi", NA, "other", "Váy dài hoa được phối thêm túi và mũ", "AI_A"),
    6162: row("hat", "good", "yes", "nâu;be", NA, NA, "trơn", NA, NA, NA, NA, "Hai biến thể màu của mũ bucket", "AI_A"),
    4840: row("jacket", "good", "yes", "xanh_navy", "tay_dài", "cổ_tròn", "trơn", "suông", "other", NA, NA, "Áo khoác cổ tròn có hàng khuy vàng; ảnh phối thêm túi", "AI_A"),
    970: row("pants", "usable", "yes", "xanh_dương;xám", NA, NA, "trơn", "rộng", NA, "ống_rộng", "cạp_cao", "Nhiều biến thể màu của quần denim ống rộng", "AI_A"),
    12701: row("set", "usable", "yes", "vàng;trắng", "tay_dài", "cổ_tròn", "trơn", "suông", "other", "not_visible", "not_visible", "Bộ áo dài và quần; phần ống và cạp quần không hiển thị đầy đủ", "AI_A"),
    9169: row("skirt", "good", "yes", "đen", NA, NA, "trơn", "suông", "midi", NA, "cạp_cao", "Cùng một chân váy trên hai người mẫu", "AI_A"),
    11443: row("jacket", "usable", "yes", "be", "không_tay", "cổ_v", "trơn", "rộng", "other", NA, NA, "Áo ghi-lê chần bông được phối với áo len và quần", "AI_A"),
    5383: row("bag", "good", "no", "kem;nâu", NA, NA, "trơn", NA, NA, NA, NA, "Túi xách phối hai màu", "AI_A"),
    5345: row("bag", "good", "no", "trắng;đen", NA, NA, "trơn", NA, NA, NA, NA, "Túi tote có quai màu đen", "AI_A"),
    2957: row("skirt", "good", "yes", "nâu;đen", NA, NA, "trơn", "suông", "midi", NA, "cạp_cao", "Hai biến thể màu của chân váy chữ A", "AI_A"),
    5995: row("hat", "good", "no", "be;đen", NA, NA, "trơn", NA, NA, NA, NA, "Mũ cói có dải màu đen", "AI_A"),
    2094: row("jacket", "good", "yes", "đen", "tay_dài", "other", "in_hình", "rộng", "other", NA, NA, "Áo hoodie có họa tiết vẩy màu; ảnh bày thêm giày và phụ kiện", "AI_A"),
    12076: row("pants", "usable", "yes", "xám", NA, NA, "in_hình", "rộng", NA, "ống_rộng", "cạp_cao", "Quần denim có các hình thêu nhỏ; ảnh phối thêm áo", "AI_A"),
    13095: row("set", "good", "yes", "xanh_navy", "tay_dài", "other", "trơn", "suông", "other", "ống_suông", "unknown", "Bộ áo cổ đứng và quần dài", "AI_A"),
    9122: row("skirt", "usable", "yes", "đen", NA, NA, "trơn", "suông", "maxi", NA, "cạp_cao", "Chân váy dài được phối với áo trắng", "AI_A"),
    11094: row("jacket", "good", "yes", "unknown", "không_tay", "cổ_bẻ", "trơn", "rộng", "other", NA, NA, "Nhiều biến thể màu của áo ghi-lê lót lông", "AI_A"),
}


def validate_annotations():
    for product_id, values in ANNOTATIONS.items():
        for column, allowed in ALLOWED.items():
            value = values[column]
            if value in (None, ""):
                raise ValueError(f"Blank required label: product_id={product_id}, column={column}")
            invalid = set(str(value).split(";")) - allowed
            if invalid:
                raise ValueError(
                    f"Invalid value(s) {sorted(invalid)}: product_id={product_id}, column={column}"
                )


def main():
    path = Path("outputs/prepared/blind_505_before_AI_A.xlsx")
    backup = path.with_name("blind_505_before_AI_A_original.xlsx")
    validate_annotations()
    workbook = load_workbook(path)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    expected = ["product_id", "split", "image_path", *COLUMNS]
    if headers != expected:
        raise ValueError(f"Unexpected workbook columns: {headers}")
    product_rows = {int(sheet.cell(r, 1).value): r for r in range(2, sheet.max_row + 1)}
    if len(product_rows) != 50 or set(product_rows) != set(ANNOTATIONS):
        missing = sorted(set(product_rows) - set(ANNOTATIONS))
        extra = sorted(set(ANNOTATIONS) - set(product_rows))
        raise ValueError(f"Expected the reviewed 50 rows; missing={missing}, extra={extra}")
    if not backup.exists():
        shutil.copy2(path, backup)
    column_index = {name: index + 1 for index, name in enumerate(headers)}
    for product_id, values in ANNOTATIONS.items():
        excel_row = product_rows[product_id]
        for column, value in values.items():
            sheet.cell(excel_row, column_index[column]).value = value
    temporary = path.with_name("blind_505_before_AI_A.__tmp__.xlsx")
    workbook.save(temporary)
    output = path
    try:
        temporary.replace(path)
    except PermissionError:
        output = path.with_name("blind_505_after_AI_A.xlsx")
        temporary.replace(output)
        print(f"Source workbook is locked; wrote the completed workbook to {output}")
    print(f"Filled {len(ANNOTATIONS)} rows in {output}")
    print(f"Backup: {backup}")


if __name__ == "__main__":
    main()
