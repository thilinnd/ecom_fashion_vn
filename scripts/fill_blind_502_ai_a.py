"""Fill blind_502_before_AI_A.xlsx from an image-only AI_A review."""
from pathlib import Path
import shutil

from openpyxl import load_workbook


COLUMNS = [
    "category_gold", "image_quality", "multiple_products", "color_gold",
    "sleeve_gold", "neckline_gold", "pattern_gold", "fit_gold", "length_gold",
    "pant_cut_gold", "waist_gold", "notes", "annotator",
]
NA = "not_applicable"


def row(*values):
    if len(values) != len(COLUMNS):
        raise ValueError(f"Expected {len(COLUMNS)} values, got {len(values)}")
    return dict(zip(COLUMNS, values))


ANNOTATIONS = {
    40: row("unknown", "usable", "yes", "xanh_navy;be;đỏ;xám;trắng;đen", "unknown", "unknown", "trơn", "unknown", "not_visible", "unknown", "unknown", "Sản phẩm được cuộn nên không xác định chắc loại và cấu trúc", "AI_A"),
    5307: row("bag", "good", "no", "be", NA, NA, "in_hình", NA, NA, NA, NA, "Túi có chữ in ở mặt trước", "AI_A"),
    3641: row("dress", "good", "no", "nâu", "không_tay", "cổ_tròn", "trơn", "suông", "midi", NA, "other", "", "AI_A"),
    6066: row("other", "usable", "yes", "unknown", NA, NA, "trơn", NA, NA, NA, NA, "Mũ bảo hiểm trẻ em ngoài taxonomy; ảnh có nhiều biến thể màu", "AI_A"),
    2631: row("jacket", "usable", "yes", "xanh_navy", "tay_dài", "cổ_lọ", "trơn", "rộng", "other", NA, NA, "Ảnh phối áo khoác với áo và quần khác", "AI_A"),
    4184: row("set", "usable", "yes", "unknown", "tay_ngắn", "cổ_tròn", "trơn", "ôm", "croptop", "ống_bó", "cạp_cao", "Nhiều biến thể màu của bộ đồ thể thao", "AI_A"),
    4477: row("set", "good", "yes", "unknown", "tay_ngắn", "cổ_tròn", "in_hình", "suông", "other", NA, "unknown", "Nhiều biến thể bộ quần áo bóng đá trẻ em", "AI_A"),
    8899: row("set", "good", "yes", "trắng;xám", "không_tay", "cổ_tròn", "trơn", "ôm", "mini", NA, "cạp_cao", "Ảnh phối áo croptop và chân váy ngắn", "AI_A"),
    9795: row("top", "good", "no", "đen", "tay_dài", "cổ_v", "trơn", "ôm", "other", NA, NA, "", "AI_A"),
    5589: row("bag", "usable", "yes", "nâu", NA, NA, "trơn", NA, NA, NA, NA, "Cùng một túi ở nhiều khung hình", "AI_A"),
    5067: row("bag", "good", "no", "trắng", NA, NA, "trơn", NA, NA, NA, NA, "", "AI_A"),
    3607: row("set", "usable", "yes", "xám;nâu;trắng", "tay_dài", "cổ_bẻ", "trơn", "suông", "mini", NA, "other", "Bộ váy yếm phối áo ren; có hai biến thể màu", "AI_A"),
    6334: row("hat", "good", "yes", "trắng", NA, NA, "in_hình", NA, NA, NA, NA, "Cùng một mũ ở mặt trước và sau", "AI_A"),
    11148: row("jacket", "good", "yes", "be", "tay_dài", "cổ_bẻ", "trơn", "rộng", "other", NA, NA, "Ảnh phối thêm áo croptop và quần ngắn", "AI_A"),
    1389: row("pants", "good", "yes", "be", NA, NA, "trơn", "rộng", NA, "ống_rộng", "cạp_cao", "Cùng một kiểu quần ở nhiều khung hình; ảnh phối thêm áo", "AI_A"),
    3861: row("set", "good", "yes", "xanh_dương", "tay_ngắn", "cổ_v", "in_hình", "suông", "other", NA, "unknown", "Bộ áo bóng đá và quần ngắn", "AI_A"),
    8585: row("set", "usable", "yes", "trắng;xanh_dương", "tay_ngắn", "cổ_v", "hoa", "suông", "midi", NA, "cạp_cao", "Ảnh phối áo trắng và chân váy hoa", "AI_A"),
    633: row("top", "good", "yes", "đen;trắng", "tay_dài", "cổ_bẻ", "trơn", "suông", "other", NA, NA, "Hai biến thể màu của áo sơ mi", "AI_A"),
    11267: row("other", "usable", "yes", "unknown", NA, NA, "trơn", NA, NA, NA, NA, "Khăn choàng lông ngoài taxonomy; nhiều biến thể màu", "AI_A"),
    5371: row("shoe", "usable", "yes", "unknown", NA, NA, "in_hình", NA, NA, NA, NA, "Nhiều biến thể dép có họa tiết monogram", "AI_A"),
    7465: row("dress", "usable", "no", "xám", "không_tay", "hai_dây", "hoa", "ôm", "maxi", NA, "other", "Váy dạ hội có bề mặt hoa ánh kim", "AI_A"),
    6245: row("hat", "good", "yes", "đỏ", NA, NA, "trơn", NA, NA, NA, NA, "Cùng một mũ beret ở nhiều góc nhìn", "AI_A"),
    10978: row("jacket", "good", "yes", "đen", "tay_dài", "other", "trơn", "rộng", "other", NA, NA, "Áo parka có mũ; cùng sản phẩm ở ảnh mẫu và ảnh minh họa", "AI_A"),
    1264: row("pants", "good", "yes", "xanh_navy", NA, NA, "trơn", "suông", NA, "ống_đứng", "not_visible", "Quần được trình bày trên hai người mẫu", "AI_A"),
    4501: row("set", "good", "yes", "unknown", "không_tay", "cổ_tròn", "in_hình", "suông", "other", NA, "unknown", "Hai biến thể màu của bộ đồ bóng rổ trẻ em", "AI_A"),
    8758: row("set", "usable", "yes", "đen;trắng", "tay_dài", "cổ_tròn", "trơn", "suông", "maxi", NA, "cạp_cao", "Ảnh phối áo đen với chân váy dài nhiều tầng", "AI_A"),
    10378: row("top", "good", "no", "nâu", "tay_ngắn", "trễ_vai", "trơn", "suông", "other", NA, NA, "Áo có phần bèo và khoét vai", "AI_A"),
    12685: row("set", "usable", "yes", "vàng;đen", "tay_dài", "cổ_tròn", "trơn", "suông", "other", "ống_rộng", "not_visible", "Ảnh phối áo vàng và quần đen ống rộng", "AI_A"),
    5626: row("bag", "good", "no", "xám;nâu;đỏ", NA, NA, "in_hình", NA, NA, NA, NA, "Ba lô có hình cú mèo", "AI_A"),
    3748: row("dress", "usable", "yes", "kem", "tay_ngắn", "cổ_v", "hoa", "suông", "maxi", NA, "other", "Cùng một váy ở hai khung hình", "AI_A"),
    5844: row("hat", "good", "yes", "unknown", NA, NA, "trơn", NA, NA, NA, NA, "Nhiều biến thể màu của mũ bucket", "AI_A"),
    11604: row("jacket", "usable", "yes", "xanh_dương", "không_tay", "cổ_bẻ", "trơn", "rộng", "other", NA, NA, "Áo khoác denim không tay; ảnh phối thêm áo và quần", "AI_A"),
    11949: row("set", "good", "yes", "đen", "tay_dài", "trễ_vai", "trơn", "ôm", "croptop", "ống_rộng", "cạp_cao", "Bộ áo croptop và quần dài bèo tầng", "AI_A"),
    3866: row("jacket", "usable", "yes", "đỏ", "tay_dài", "cổ_lọ", "in_hình", "suông", "other", NA, NA, "Ảnh phối thêm áo và quần đen", "AI_A"),
    9012: row("set", "usable", "yes", "nâu", "tay_dài", "cổ_v", "caro", "suông", "midi", NA, "cạp_cao", "Ảnh phối cardigan với chân váy caro", "AI_A"),
    10133: row("top", "usable", "yes", "trắng;nâu", "tay_ngắn", "cổ_tròn", "kẻ_sọc", "suông", "other", NA, NA, "Ảnh phối thêm mũ lưỡi trai", "AI_A"),
    1195: row("pants", "usable", "yes", "xanh_dương;đen", NA, NA, "trơn", "rộng", NA, NA, "cạp_cao", "Quần denim ngắn có nhiều biến thể màu; pant cut không áp dụng", "AI_A"),
    5435: row("bag", "good", "no", "nâu", NA, NA, "trơn", NA, NA, NA, NA, "Túi có móc trang trí", "AI_A"),
    6960: row("dress", "usable", "yes", "xanh_dương;trắng", "không_tay", "cổ_yếm", "trơn", "suông", "maxi", NA, "other", "Ảnh có thêm một biến thể váy trắng ở hậu cảnh", "AI_A"),
    12903: row("set", "usable", "yes", "xanh_lá;đỏ", "tay_dài", "cổ_bẻ", "trơn", "suông", "other", "ống_bó", "unknown", "Hai trẻ mặc cùng bộ áo và quần", "AI_A"),
    2028: row("jacket", "good", "yes", "be", "tay_lỡ", "cổ_bẻ", "trơn", "rộng", "other", NA, NA, "Ảnh có minh họa thêm các biến thể màu", "AI_A"),
    1010: row("pants", "usable", "yes", "xanh_dương", NA, NA, "trơn", "rộng", NA, "ống_rộng", "cạp_cao", "Quần denim trẻ em; ảnh phối thêm áo", "AI_A"),
    3917: row("set", "good", "yes", "unknown", "tay_dài", "cổ_lọ", "trơn", "suông", "other", "ống_suông", "unknown", "Ba biến thể màu của bộ đồ thể thao", "AI_A"),
    10313: row("top", "usable", "yes", "unknown", "không_tay", "hai_dây", "trơn", "ôm", "croptop", NA, NA, "Nhiều biến thể màu của áo nhún bèo", "AI_A"),
    396: row("set", "good", "yes", "hồng", "tay_ngắn", "cổ_tròn", "in_hình", "rộng", "other", NA, "unknown", "Bộ áo thun và quần ngắn trẻ em", "AI_A"),
    5087: row("bag", "usable", "no", "xanh_dương", NA, NA, "trơn", NA, NA, NA, NA, "Ảnh sản phẩm có độ phân giải thấp", "AI_A"),
    6836: row("dress", "good", "no", "nâu", "tay_ngắn", "một_vai", "trơn", "ôm", "mini", NA, "other", "Váy bất đối xứng một vai", "AI_A"),
    5780: row("hat", "good", "yes", "xanh_navy", NA, NA, "in_hình", NA, NA, NA, NA, "Cùng một mũ ở mặt trước và sau", "AI_A"),
    1794: row("jacket", "good", "yes", "nâu", "tay_dài", "cổ_bẻ", "trơn", "suông", "other", NA, NA, "Ảnh phối thêm trang phục và túi màu đen", "AI_A"),
    12196: row("set", "usable", "yes", "đen;xám", "không_tay", "cổ_lọ", "other", "rộng", "croptop", "ống_rộng", "cạp_cao", "Bộ áo croptop và quần denim rách; chi tiết rách ngoài vocabulary pattern", "AI_A"),
}


def main():
    path = Path("outputs/prepared/blind_502_before_AI_A.xlsx")
    backup = path.with_name("blind_502_before_AI_A_original.xlsx")
    workbook = load_workbook(path)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    expected = ["product_id", "split", "image_path", *COLUMNS]
    if headers != expected:
        raise ValueError(f"Unexpected workbook columns: {headers}")
    product_rows = {int(sheet.cell(r, 1).value): r for r in range(2, sheet.max_row + 1)}
    if len(product_rows) != 50 or set(product_rows) != set(ANNOTATIONS):
        missing = sorted(set(product_rows) - set(ANNOTATIONS))
        extra = sorted(set(ANNOTATIONS) - set(product_rows))
        raise ValueError(f"Expected the reviewed 50 rows; missing={missing}, extra={extra}")
    if not backup.exists():
        shutil.copy2(path, backup)
    column_index = {name: index + 1 for index, name in enumerate(headers)}
    for product_id, values in ANNOTATIONS.items():
        excel_row = product_rows[product_id]
        for column, value in values.items():
            sheet.cell(excel_row, column_index[column]).value = value
    temporary = path.with_name("blind_502_before_AI_A.__tmp__.xlsx")
    workbook.save(temporary)
    temporary.replace(path)
    print(f"Filled {len(ANNOTATIONS)} rows in {path}")
    print(f"Backup: {backup}")


if __name__ == "__main__":
    main()
