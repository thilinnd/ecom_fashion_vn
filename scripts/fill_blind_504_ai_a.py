"""Create and fill blind_504_before_AI_A.xlsx from an image-only AI_A review."""
from pathlib import Path

from openpyxl import load_workbook


COLUMNS = [
    "category_gold", "image_quality", "multiple_products", "color_gold",
    "sleeve_gold", "neckline_gold", "pattern_gold", "fit_gold", "length_gold",
    "pant_cut_gold", "waist_gold", "notes", "annotator",
]
NA = "not_applicable"
SPECIAL = {"unknown", "not_visible", NA, "other"}
ALLOWED = {
    "category_gold": {"top", "dress", "pants", "skirt", "jacket", "bag", "hat", "shoe", "set", "fabric", "other", "unknown"},
    "image_quality": {"good", "usable", "poor", "unusable"},
    "multiple_products": {"yes", "no", "unknown"},
    "color_gold": {"đen", "trắng", "đỏ", "hồng", "xanh_dương", "xanh_lá", "xanh_navy", "vàng", "tím", "nâu", "xám", "be", "kem"} | SPECIAL,
    "sleeve_gold": {"tay_dài", "tay_ngắn", "tay_lỡ", "không_tay", "tay_phồng", "tay_loe", "tay_raglan"} | SPECIAL,
    "neckline_gold": {"cổ_tròn", "cổ_v", "cổ_vuông", "cổ_lọ", "cổ_yếm", "cổ_đổ", "cổ_bẻ", "trễ_vai", "hai_dây", "một_vai"} | SPECIAL,
    "pattern_gold": {"trơn", "kẻ_sọc", "caro", "chấm_bi", "hoa", "in_hình", "loang", "da_báo"} | SPECIAL,
    "fit_gold": {"rộng", "ôm", "suông"} | SPECIAL,
    "length_gold": {"maxi", "midi", "mini", "croptop"} | SPECIAL,
    "pant_cut_gold": {"ống_rộng", "ống_suông", "ống_loe", "ống_đứng", "ống_bó"} | SPECIAL,
    "waist_gold": {"cạp_cao", "cạp_thấp"} | SPECIAL,
    "annotator": {"AI_A"},
}


def row(*values):
    if len(values) != len(COLUMNS):
        raise ValueError(f"Expected {len(COLUMNS)} values, got {len(values)}")
    return dict(zip(COLUMNS, values))


ANNOTATIONS = {
    4497: row("set", "good", "yes", "đen", "tay_dài", "cổ_tròn", "trơn", "ôm", "other", "ống_bó", "cạp_cao", "Bộ áo và quần giữ nhiệt trẻ em", "AI_A"),
    7440: row("dress", "good", "no", "hồng", "tay_dài", "cổ_tròn", "trơn", "suông", "mini", NA, "other", "Váy có nhiều tầng bèo trang trí", "AI_A"),
    9390: row("set", "good", "yes", "xanh_navy", "tay_lỡ", "cổ_bẻ", "trơn", "rộng", "other", "ống_rộng", "not_visible", "Bộ áo và quần dáng rộng", "AI_A"),
    11276: row("other", "usable", "yes", "xám", NA, NA, "trơn", NA, NA, NA, NA, "Khăn choàng lông ngoài taxonomy; ảnh phối nhiều trang phục", "AI_A"),
    5538: row("bag", "good", "no", "trắng", NA, NA, "other", NA, NA, NA, NA, "Túi có bề mặt chần quả trám ngoài vocabulary pattern", "AI_A"),
    3062: row("dress", "good", "no", "trắng", "tay_ngắn", "cổ_vuông", "trơn", "rộng", "mini", NA, "other", "Váy babydoll dáng rộng", "AI_A"),
    5850: row("hat", "good", "yes", "unknown", NA, NA, "trơn", NA, NA, NA, NA, "Nhiều biến thể màu của mũ bucket", "AI_A"),
    11548: row("jacket", "good", "yes", "xám", "tay_dài", "other", "trơn", "rộng", "other", NA, NA, "Cùng một áo khoác có mũ ở ảnh sản phẩm và trên người mẫu", "AI_A"),
    925: row("pants", "good", "yes", "xanh_dương", NA, NA, "trơn", "suông", NA, "ống_suông", "cạp_cao", "Quần denim được chụp trên người mẫu cùng áo phối", "AI_A"),
    4179: row("set", "good", "yes", "xanh_navy", "không_tay", "cổ_tròn", "trơn", "ôm", "croptop", NA, "cạp_cao", "Bộ đồ thể thao gồm áo bra và quần ngắn", "AI_A"),
    8682: row("set", "usable", "yes", "trắng;nâu", "tay_dài", "trễ_vai", "trơn", "suông", "mini", NA, "cạp_cao", "Hai cách phối áo trắng với chân váy ngắn màu nâu", "AI_A"),
    726: row("top", "good", "yes", "vàng", "tay_dài", "trễ_vai", "trơn", "ôm", "other", NA, NA, "Ảnh phối áo với quần denim", "AI_A"),
    3962: row("top", "good", "no", "trắng", "tay_ngắn", "cổ_bẻ", "trơn", "suông", "other", NA, NA, "Áo polo được chụp từ phía sau", "AI_A"),
    5156: row("bag", "good", "no", "đen", NA, NA, "trơn", NA, NA, NA, NA, "", "AI_A"),
    8462: row("skirt", "good", "no", "đỏ", NA, NA, "trơn", "suông", "mini", NA, "cạp_cao", "Chân váy xếp ly", "AI_A"),
    6168: row("hat", "usable", "yes", "xanh_navy;be", NA, NA, "trơn", NA, NA, NA, NA, "Hai biến thể màu của mũ bucket trẻ em", "AI_A"),
    2022: row("jacket", "good", "yes", "unknown", "tay_dài", "cổ_bẻ", "trơn", "suông", "other", NA, NA, "Nhiều biến thể màu của áo khoác", "AI_A"),
    940: row("pants", "good", "no", "xanh_dương", NA, NA, "trơn", "suông", NA, "ống_đứng", "cạp_cao", "Quần denim", "AI_A"),
    3854: row("top", "good", "no", "xanh_dương;đỏ;trắng", "không_tay", "cổ_v", "in_hình", "suông", "other", NA, NA, "Áo thể thao có chữ và số in", "AI_A"),
    8010: row("skirt", "good", "no", "hồng", NA, NA, "trơn", "suông", "maxi", NA, "cạp_cao", "Chân váy dài nhiều tầng", "AI_A"),
    10964: row("jacket", "good", "no", "đen", "tay_dài", "other", "trơn", "rộng", "other", NA, NA, "Áo khoác có mũ và lớp lót lông", "AI_A"),
    5559: row("bag", "good", "no", "đen", NA, NA, "trơn", NA, NA, NA, NA, "Túi đeo chéo nhiều ngăn", "AI_A"),
    5272: row("bag", "good", "yes", "xám", NA, NA, "trơn", NA, NA, NA, NA, "Hai túi đeo chéo trên người mẫu; trang phục không phải sản phẩm đích", "AI_A"),
    3443: row("dress", "good", "no", "be", "tay_ngắn", "cổ_bẻ", "trơn", "ôm", "midi", NA, "other", "Váy sơ mi có đai eo", "AI_A"),
    6375: row("hat", "good", "no", "xanh_navy", NA, NA, "in_hình", NA, NA, NA, NA, "Mũ lưỡi trai có chữ thêu ở mặt trước", "AI_A"),
    11121: row("jacket", "usable", "yes", "xanh_dương", "tay_dài", "cổ_bẻ", "trơn", "rộng", "other", NA, NA, "Gán thuộc tính cho blazer; ảnh phối thêm áo và quần ngắn", "AI_A"),
    11974: row("pants", "good", "yes", "unknown", NA, NA, "trơn", "suông", NA, "ống_đứng", "cạp_cao", "Nhiều biến thể màu của quần", "AI_A"),
    4490: row("set", "good", "yes", "đỏ;trắng", "tay_ngắn", "cổ_tròn", "in_hình", "suông", "other", NA, "unknown", "Bộ áo bóng đá và quần ngắn", "AI_A"),
    8658: row("set", "usable", "yes", "đen;trắng", "tay_dài", "cổ_tròn", "chấm_bi", "suông", "midi", NA, "cạp_cao", "Ảnh phối áo đen với chân váy chấm bi", "AI_A"),
    752: row("top", "good", "no", "trắng", "tay_dài", "cổ_bẻ", "trơn", "rộng", "other", NA, NA, "Áo sơ mi dáng rộng", "AI_A"),
    12421: row("dress", "good", "yes", "xanh_dương", "không_tay", "cổ_tròn", "trơn", "suông", "mini", NA, "other", "Cùng một váy ở mặt trước và sau", "AI_A"),
    5301: row("bag", "good", "no", "đen", NA, NA, "trơn", NA, NA, NA, NA, "", "AI_A"),
    3238: row("dress", "good", "no", "kem", "tay_ngắn", "cổ_vuông", "trơn", "ôm", "maxi", NA, "other", "Váy dài có lớp voan phủ", "AI_A"),
    6038: row("hat", "usable", "yes", "đỏ", NA, NA, "in_hình", NA, NA, NA, NA, "Nhiều mũ trẻ em có hình ngôi sao", "AI_A"),
    2431: row("set", "usable", "yes", "trắng;xanh_dương", "tay_dài", "cổ_bẻ", "caro", "suông", "mini", NA, "cạp_cao", "Bộ áo và chân váy caro", "AI_A"),
    1400: row("pants", "good", "yes", "đen;be", NA, NA, "trơn", "rộng", NA, "ống_rộng", "cạp_cao", "Hai biến thể màu của quần ống rộng", "AI_A"),
    3887: row("set", "good", "yes", "đỏ", "tay_ngắn", "cổ_v", "in_hình", "suông", "other", NA, "unknown", "Bộ áo bóng đá và quần ngắn", "AI_A"),
    8883: row("skirt", "usable", "yes", "trắng", NA, NA, "trơn", "suông", "mini", NA, "cạp_cao", "Cùng một chân váy xếp ly ở nhiều khung hình", "AI_A"),
    2324: row("jacket", "good", "yes", "xám;đen", "tay_dài", "other", "trơn", "rộng", "other", NA, NA, "Áo phao được trình bày ở hai góc nhìn", "AI_A"),
    4275: row("set", "usable", "yes", "unknown", "tay_ngắn", "cổ_tròn", "trơn", "rộng", "croptop", "ống_bó", "cạp_cao", "Nhiều biến thể màu của bộ áo croptop và quần jogger", "AI_A"),
    5611: row("bag", "good", "no", "nâu", NA, NA, "trơn", NA, NA, NA, NA, "Túi rút bằng vải gân", "AI_A"),
    3205: row("dress", "usable", "no", "xanh_dương", "tay_lỡ", "cổ_bẻ", "trơn", "rộng", "midi", NA, "other", "Váy denim dáng rộng", "AI_A"),
    2290: row("jacket", "good", "no", "nâu;kem", "tay_dài", "other", "trơn", "rộng", "other", NA, NA, "Áo hoodie phối hai màu", "AI_A"),
    12856: row("set", "usable", "yes", "xanh_dương;vàng;hồng", "tay_lỡ", "cổ_tròn", "trơn;hoa", "suông", "other", "ống_rộng", "not_visible", "Nhiều biến thể của bộ áo dài truyền thống và quần", "AI_A"),
    12399: row("set", "good", "yes", "unknown", "không_tay", "cổ_vuông", "caro", "suông", "other", NA, "cạp_cao", "Nhiều biến thể màu của bộ áo và quần ngắn caro", "AI_A"),
    8619: row("set", "good", "yes", "trắng;đen", "tay_dài", "cổ_tròn", "trơn", "ôm", "midi", NA, "cạp_cao", "Ảnh phối áo trắng với chân váy bút chì", "AI_A"),
    11798: row("top", "good", "yes", "unknown", "không_tay", "hai_dây", "trơn", "ôm", "croptop", NA, NA, "Nhiều biến thể màu của áo bra thể thao", "AI_A"),
    445: row("jacket", "usable", "yes", "unknown", "tay_dài", "cổ_bẻ", "trơn", "rộng", "other", NA, NA, "Nhiều biến thể màu của áo khoác sơ mi", "AI_A"),
    5025: row("bag", "good", "no", "kem", NA, NA, "other", NA, NA, NA, NA, "Túi có bề mặt chần ngoài vocabulary pattern", "AI_A"),
    7268: row("dress", "good", "no", "be", "không_tay", "một_vai", "trơn", "ôm", "maxi", NA, "other", "Váy dài có phần thân trên bất đối xứng", "AI_A"),
}


def validate_annotations():
    for product_id, values in ANNOTATIONS.items():
        for column, allowed in ALLOWED.items():
            value = values[column]
            if value in (None, ""):
                raise ValueError(f"Blank required label: product_id={product_id}, column={column}")
            invalid = set(str(value).split(";")) - allowed
            if invalid:
                raise ValueError(
                    f"Invalid value(s) {sorted(invalid)}: product_id={product_id}, column={column}"
                )


def main():
    source = Path("outputs/prepared/blind_504_before_AI_A_original.xlsx")
    target = Path("outputs/prepared/blind_504_before_AI_A.xlsx")
    validate_annotations()
    workbook = load_workbook(source)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    expected = ["product_id", "split", "image_path", *COLUMNS]
    if headers != expected:
        raise ValueError(f"Unexpected workbook columns: {headers}")
    product_rows = {int(sheet.cell(r, 1).value): r for r in range(2, sheet.max_row + 1)}
    if len(product_rows) != 50 or set(product_rows) != set(ANNOTATIONS):
        missing = sorted(set(product_rows) - set(ANNOTATIONS))
        extra = sorted(set(ANNOTATIONS) - set(product_rows))
        raise ValueError(f"Expected the reviewed 50 rows; missing={missing}, extra={extra}")
    column_index = {name: index + 1 for index, name in enumerate(headers)}
    for product_id, values in ANNOTATIONS.items():
        excel_row = product_rows[product_id]
        for column, value in values.items():
            sheet.cell(excel_row, column_index[column]).value = value
    temporary = target.with_name("blind_504_before_AI_A.__tmp__.xlsx")
    workbook.save(temporary)
    temporary.replace(target)
    print(f"Filled {len(ANNOTATIONS)} rows in {target}")
    print(f"Unchanged source: {source}")


if __name__ == "__main__":
    main()
