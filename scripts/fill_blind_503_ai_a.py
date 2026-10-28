"""Fill blind_503_before_AI_A.xlsx from an image-only AI_A review."""
from pathlib import Path
import shutil

from openpyxl import load_workbook


COLUMNS = [
    "category_gold", "image_quality", "multiple_products", "color_gold",
    "sleeve_gold", "neckline_gold", "pattern_gold", "fit_gold", "length_gold",
    "pant_cut_gold", "waist_gold", "notes", "annotator",
]
NA = "not_applicable"


def row(*values):
    if len(values) != len(COLUMNS):
        raise ValueError(f"Expected {len(COLUMNS)} values, got {len(values)}")
    return dict(zip(COLUMNS, values))


ANNOTATIONS = {
    5704: row("bag", "usable", "yes", "unknown", NA, NA, "in_hình", NA, NA, NA, NA, "Nhiều biến thể màu của túi đeo hông", "AI_A"),
    7689: row("set", "usable", "yes", "đỏ;đen", "tay_dài", "cổ_tròn", "trơn", "ôm", "mini;midi", NA, "other", "Ảnh chứa hai thiết kế khác nhau: bộ áo khoác và váy ngắn cùng váy liền midi", "AI_A"),
    574: row("jacket", "good", "yes", "unknown", "tay_dài", "cổ_bẻ", "trơn", "rộng", "other", NA, NA, "Nhiều biến thể màu của áo khoác sơ mi", "AI_A"),
    5979: row("hat", "good", "no", "đen", NA, NA, "trơn", NA, NA, NA, NA, "", "AI_A"),
    4955: row("jacket", "good", "no", "trắng;đen", "tay_dài", "cổ_bẻ", "kẻ_sọc", "ôm", "other", NA, NA, "", "AI_A"),
    4249: row("set", "usable", "yes", "đen", "không_tay", "cổ_tròn", "trơn", "ôm", NA, "ống_loe", "cạp_cao", "Trang phục liền quần ngoài taxonomy; ảnh có nhiều góc nhìn và biến thể màu", "AI_A"),
    4515: row("set", "good", "yes", "đỏ;đen", "tay_ngắn", "cổ_tròn", "in_hình", "rộng", "other", NA, "unknown", "Bộ áo thun và quần ngắn", "AI_A"),
    8710: row("skirt", "good", "yes", "nâu", NA, NA, "trơn", "suông", "midi", NA, "cạp_cao", "Ảnh phối chân váy với áo trắng", "AI_A"),
    2496: row("top", "good", "yes", "đen;trắng", "tay_dài", "cổ_bẻ", "in_hình", "rộng", "other", NA, NA, "Hai biến thể màu của áo trẻ em", "AI_A"),
    5685: row("bag", "good", "yes", "unknown", NA, NA, "in_hình", NA, NA, NA, NA, "Nhiều biến thể màu và hình in của túi trẻ em", "AI_A"),
    5418: row("bag", "good", "yes", "kem", NA, NA, "trơn", NA, NA, NA, NA, "Nhiều kích thước túi vải trong cùng ảnh", "AI_A"),
    7978: row("dress", "good", "no", "đỏ", "không_tay", "một_vai", "trơn", "ôm", "maxi", NA, "other", "Váy dạ hội có chi tiết hoa lớn ở một vai", "AI_A"),
    6125: row("hat", "good", "yes", "unknown", NA, NA, "in_hình", NA, NA, NA, NA, "Nhiều kiểu và màu mũ lưỡi trai", "AI_A"),
    4789: row("jacket", "good", "yes", "kem", "tay_dài", "cổ_bẻ", "trơn", "ôm", "other", NA, NA, "Gán thuộc tính cho blazer; ảnh phối thêm váy và túi", "AI_A"),
    1517: row("set", "usable", "yes", "xanh_navy;xám", "tay_dài", "cổ_bẻ", "trơn", "suông", "other", "ống_đứng", "cạp_cao", "Bộ áo sơ mi và quần công sở", "AI_A"),
    9291: row("top", "usable", "yes", "hồng", "không_tay", "cổ_yếm", "trơn", "suông", "other", NA, NA, "Ảnh phối thêm chân váy denim", "AI_A"),
    8045: row("skirt", "good", "yes", "xám", NA, NA, "trơn", "suông", "mini", NA, "cạp_cao", "Cùng một chân váy ở ảnh sản phẩm và trên người mẫu", "AI_A"),
    9841: row("top", "good", "yes", "unknown", "tay_ngắn", "cổ_tròn", "in_hình", "rộng", "other", NA, NA, "Nhiều biến thể màu của áo thun", "AI_A"),
    12686: row("top", "good", "no", "hồng", "tay_lỡ", "cổ_tròn", "hoa", "ôm", "other", NA, NA, "Bề mặt vải có họa tiết hoa chìm", "AI_A"),
    5375: row("bag", "good", "yes", "unknown", NA, NA, "trơn", NA, NA, NA, NA, "Nhiều biến thể màu của túi đeo điện thoại", "AI_A"),
    7437: row("dress", "good", "no", "đỏ", "tay_ngắn", "cổ_tròn", "trơn", "ôm", "mini", NA, "other", "Tay phồng ngắn", "AI_A"),
    6364: row("hat", "good", "no", "be", NA, NA, "trơn", NA, NA, NA, NA, "", "AI_A"),
    10394: row("top", "good", "no", "hồng", "không_tay", "hai_dây", "trơn", "suông", "other", NA, NA, "", "AI_A"),
    1280: row("set", "usable", "yes", "hồng;trắng", "tay_ngắn", "cổ_bẻ", "trơn", "suông", "other", "ống_đứng", "not_visible", "Ảnh phối áo polo và quần trắng", "AI_A"),
    3882: row("set", "usable", "yes", "unknown", "tay_ngắn", "cổ_v", "in_hình", "suông", "other", NA, "unknown", "Nhiều biến thể màu của bộ quần áo bóng đá", "AI_A"),
    8344: row("skirt", "usable", "no", "đen", NA, NA, "hoa", "ôm", "midi", NA, "not_visible", "Ảnh cắt cận; phần ren hoa ở đường xẻ được quan sát rõ", "AI_A"),
    131: row("top", "good", "no", "đen", "tay_ngắn", "cổ_tròn", "in_hình", "suông", "other", NA, NA, "", "AI_A"),
    12905: row("set", "usable", "yes", "kem;hồng", "tay_dài", "cổ_tròn", "hoa", "suông", "other", "ống_rộng", "not_visible", "Bộ trang phục truyền thống gồm áo và quần", "AI_A"),
    5247: row("bag", "good", "no", "xanh_lá", NA, NA, "trơn", NA, NA, NA, NA, "", "AI_A"),
    7976: row("dress", "good", "no", "be", "không_tay", "hai_dây", "trơn", "ôm", "maxi", NA, "other", "", "AI_A"),
    6254: row("hat", "good", "no", "trắng", NA, NA, "trơn", NA, NA, NA, NA, "", "AI_A"),
    1850: row("jacket", "good", "yes", "trắng;đen", "tay_dài", "cổ_bẻ", "trơn", "ôm", "other", NA, NA, "Cùng một áo khoác ở mặt trước và sau", "AI_A"),
    12257: row("pants", "good", "yes", "nâu", NA, NA, "trơn", "rộng", NA, "ống_rộng", "cạp_cao", "Quần được trình bày trên người mẫu và ảnh sản phẩm; có các chấm màu biến thể", "AI_A"),
    12740: row("set", "usable", "yes", "nâu;xanh_lá", "tay_dài", "cổ_bẻ", "trơn", "suông", "other", "ống_rộng", "unknown", "Hai trẻ mặc cùng bộ áo và quần", "AI_A"),
    8735: row("skirt", "good", "yes", "đen", NA, NA, "trơn", "suông", "mini", NA, "cạp_cao", "Chân váy được trình bày trên người mẫu và ảnh sản phẩm", "AI_A"),
    10426: row("top", "good", "no", "xám", "không_tay", "cổ_tròn", "trơn", "rộng", "other", NA, NA, "Áo không tay có đệm vai", "AI_A"),
    4048: row("set", "usable", "yes", "đen", "tay_ngắn", "cổ_tròn", "in_hình", "rộng", "other", NA, "unknown", "Bộ áo thun và quần ngắn", "AI_A"),
    5010: row("bag", "good", "no", "nâu", NA, NA, "trơn", NA, NA, NA, NA, "", "AI_A"),
    3135: row("set", "usable", "yes", "trắng;đen", "tay_dài", "cổ_v", "trơn", "ôm", "mini", NA, "cạp_cao", "Hai biến thể màu của bộ áo và chân váy xếp ly", "AI_A"),
    5857: row("hat", "good", "no", "đen", NA, NA, "trơn", NA, NA, NA, NA, "", "AI_A"),
    2717: row("jacket", "good", "no", "đen", "tay_dài", "cổ_lọ", "trơn", "rộng", "other", NA, NA, "", "AI_A"),
    1191: row("pants", "good", "no", "đen", NA, NA, "other", "ôm", NA, NA, "cạp_cao", "Quần denim ngắn có chi tiết rách ngoài vocabulary pattern", "AI_A"),
    13159: row("set", "good", "yes", "nâu;xanh_lá;be", "tay_ngắn", "cổ_v", "hoa", "suông", "other", "ống_suông", "unknown", "Bộ áo và quần có cùng họa tiết", "AI_A"),
    8566: row("skirt", "usable", "yes", "xanh_dương", NA, NA, "trơn", "ôm", "mini", NA, "cạp_cao", "Chân váy denim được trình bày cận cảnh và trên người mẫu", "AI_A"),
    638: row("dress", "usable", "yes", "trắng;đen", "tay_ngắn", "cổ_bẻ", "trơn", "ôm", "mini", NA, "other", "Hai biến thể màu của váy sơ mi", "AI_A"),
    5343: row("bag", "usable", "no", "đen", NA, NA, "trơn", NA, NA, NA, NA, "Túi được chụp khi đang sử dụng", "AI_A"),
    5602: row("bag", "good", "no", "be", NA, NA, "trơn", NA, NA, NA, NA, "Túi có quai xích và túi tròn nhỏ đi kèm", "AI_A"),
    3408: row("dress", "good", "no", "trắng", "không_tay", "hai_dây", "hoa", "ôm", "mini", NA, "other", "", "AI_A"),
    5771: row("hat", "good", "yes", "xanh_navy", NA, NA, "in_hình", NA, NA, NA, NA, "Cùng một mũ ở mặt trước và sau", "AI_A"),
    1587: row("pants", "usable", "yes", "xanh_navy", NA, NA, "trơn", "unknown", NA, "not_visible", "not_visible", "Hai quần được gấp nên không quan sát chắc phom ống và cạp", "AI_A"),
}


def main():
    path = Path("outputs/prepared/blind_503_before_AI_A.xlsx")
    backup = path.with_name("blind_503_before_AI_A_original.xlsx")
    workbook = load_workbook(path)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    expected = ["product_id", "split", "image_path", *COLUMNS]
    if headers != expected:
        raise ValueError(f"Unexpected workbook columns: {headers}")
    product_rows = {int(sheet.cell(r, 1).value): r for r in range(2, sheet.max_row + 1)}
    if len(product_rows) != 50 or set(product_rows) != set(ANNOTATIONS):
        missing = sorted(set(product_rows) - set(ANNOTATIONS))
        extra = sorted(set(ANNOTATIONS) - set(product_rows))
        raise ValueError(f"Expected the reviewed 50 rows; missing={missing}, extra={extra}")
    if not backup.exists():
        shutil.copy2(path, backup)
    column_index = {name: index + 1 for index, name in enumerate(headers)}
    for product_id, values in ANNOTATIONS.items():
        excel_row = product_rows[product_id]
        for column, value in values.items():
            sheet.cell(excel_row, column_index[column]).value = value
    temporary = path.with_name("blind_503_before_AI_A.__tmp__.xlsx")
    workbook.save(temporary)
    temporary.replace(path)
    print(f"Filled {len(ANNOTATIONS)} rows in {path}")
    print(f"Backup: {backup}")


if __name__ == "__main__":
    main()
