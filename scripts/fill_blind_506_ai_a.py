"""Fill blind_506_before_AI_A.xlsx from an image-only AI_A review."""
from pathlib import Path
import shutil

from openpyxl import load_workbook


COLUMNS = [
    "category_gold", "image_quality", "multiple_products", "color_gold",
    "sleeve_gold", "neckline_gold", "pattern_gold", "fit_gold", "length_gold",
    "pant_cut_gold", "waist_gold", "notes", "annotator",
]
NA = "not_applicable"
SPECIAL = {"unknown", "not_visible", NA, "other"}
ALLOWED = {
    "category_gold": {"top", "dress", "pants", "skirt", "jacket", "bag", "hat", "shoe", "set", "fabric", "other", "unknown"},
    "image_quality": {"good", "usable", "poor", "unusable"},
    "multiple_products": {"yes", "no", "unknown"},
    "color_gold": {"đen", "trắng", "đỏ", "hồng", "xanh_dương", "xanh_lá", "xanh_navy", "vàng", "tím", "nâu", "xám", "be", "kem"} | SPECIAL,
    "sleeve_gold": {"tay_dài", "tay_ngắn", "tay_lỡ", "không_tay", "tay_phồng", "tay_loe", "tay_raglan"} | SPECIAL,
    "neckline_gold": {"cổ_tròn", "cổ_v", "cổ_vuông", "cổ_lọ", "cổ_yếm", "cổ_đổ", "cổ_bẻ", "trễ_vai", "hai_dây", "một_vai"} | SPECIAL,
    "pattern_gold": {"trơn", "kẻ_sọc", "caro", "chấm_bi", "hoa", "in_hình", "loang", "da_báo"} | SPECIAL,
    "fit_gold": {"rộng", "ôm", "suông"} | SPECIAL,
    "length_gold": {"maxi", "midi", "mini", "croptop"} | SPECIAL,
    "pant_cut_gold": {"ống_rộng", "ống_suông", "ống_loe", "ống_đứng", "ống_bó"} | SPECIAL,
    "waist_gold": {"cạp_cao", "cạp_thấp"} | SPECIAL,
    "annotator": {"AI_A"},
}


def row(*values):
    if len(values) != len(COLUMNS):
        raise ValueError(f"Expected {len(COLUMNS)} values, got {len(values)}")
    return dict(zip(COLUMNS, values))


ANNOTATIONS = {
    5329: row("bag", "good", "no", "be;xám", NA, NA, "trơn", NA, NA, NA, NA, "Túi cói phối quai và viền màu xám", "AI_A"),
    5366: row("bag", "good", "yes", "unknown", NA, NA, "in_hình", NA, NA, NA, NA, "Túi phao chần bông có họa tiết nhiều màu và các biến thể minh họa", "AI_A"),
    3209: row("dress", "good", "no", "kem", "tay_ngắn", "cổ_tròn", "hoa", "ôm", "maxi", NA, "other", "Váy dài có bề mặt ren hoa và thân trên ôm", "AI_A"),
    6203: row("hat", "good", "no", "be", NA, NA, "in_hình", NA, NA, NA, NA, "Mũ lưỡi trai có hình thêu nhỏ", "AI_A"),
    2710: row("jacket", "good", "yes", "trắng", "tay_dài", "other", "trơn", "rộng", "croptop", NA, NA, "Áo khoác ngắn có mũ được phối với áo và quần đen", "AI_A"),
    11969: row("top", "good", "yes", "trắng", "không_tay", "cổ_vuông", "trơn", "ôm", "croptop", NA, NA, "Áo croptop được phối với quần ống rộng", "AI_A"),
    4009: row("top", "usable", "yes", "hồng", "tay_ngắn", "cổ_bẻ", "in_hình", "suông", "other", NA, NA, "Áo thể thao được phối với quần ngắn", "AI_A"),
    8535: row("skirt", "good", "yes", "đen;nâu;be", NA, NA, "trơn", "suông", "mini", NA, "cạp_cao", "Ba biến thể màu của chân váy xếp ly", "AI_A"),
    9461: row("top", "usable", "yes", "đen", "không_tay", "cổ_yếm", "other", "ôm", "croptop", NA, NA, "Áo croptop đính sequin được phối với quần đen", "AI_A"),
    5656: row("bag", "usable", "yes", "xám", NA, NA, "trơn", NA, NA, NA, NA, "Túi đeo chéo nhỏ được chụp trên người mẫu", "AI_A"),
    5353: row("bag", "good", "yes", "đen;xám", NA, NA, "trơn", NA, NA, NA, NA, "Túi xách được trình bày ở nhiều góc nhìn", "AI_A"),
    7826: row("dress", "good", "yes", "hồng", "tay_ngắn", "cổ_v", "hoa", "ôm", "maxi", NA, "other", "Cùng một váy dài hoa ở hai khung hình", "AI_A"),
    11425: row("jacket", "good", "yes", "xám", "tay_dài", "other", "trơn", "rộng", "other", NA, NA, "Áo phao được phối với áo, quần và túi", "AI_A"),
    13103: row("top", "good", "yes", "đỏ", "tay_ngắn", "other", "hoa", "suông", "other", NA, NA, "Áo nam cổ đứng có họa tiết thêu", "AI_A"),
    12391: row("set", "usable", "yes", "hồng", "tay_lỡ", "cổ_tròn", "hoa", "suông", "other", "not_visible", "not_visible", "Trang phục truyền thống; phần quần và cạp không hiển thị", "AI_A"),
    8077: row("skirt", "good", "yes", "unknown", NA, NA, "trơn", "suông", "mini", NA, "cạp_cao", "Chân váy chữ A có các biến thể trắng, đen và xám", "AI_A"),
    10787: row("top", "good", "yes", "trắng;đen", "không_tay", "hai_dây", "trơn", "ôm", "croptop", NA, NA, "Nhiều áo hai dây dệt móc màu trắng và đen", "AI_A"),
    12994: row("set", "good", "yes", "trắng", "tay_ngắn", "other", "in_hình", "suông", "other", "ống_suông", "unknown", "Bộ áo cổ đứng và quần dài có chi tiết trang trí", "AI_A"),
    5369: row("bag", "good", "yes", "unknown", NA, NA, "trơn", NA, NA, NA, NA, "Nhiều biến thể màu của túi đeo chéo", "AI_A"),
    7322: row("dress", "good", "yes", "xanh_navy", "tay_ngắn", "cổ_bẻ", "trơn", "suông", "midi", NA, "other", "Váy phong cách thủy thủ được phối thêm mũ", "AI_A"),
    6246: row("hat", "usable", "yes", "xám", NA, NA, "trơn", NA, NA, NA, NA, "Mũ được chụp trên người mẫu cùng trang phục khác", "AI_A"),
    10832: row("jacket", "good", "yes", "nâu", "tay_dài", "other", "trơn", "suông", "other", NA, NA, "Áo khoác da được trình bày trên người mẫu và ảnh sản phẩm", "AI_A"),
    12288: row("pants", "good", "yes", "trắng;xám", NA, NA, "trơn", "rộng", NA, "ống_bó", "cạp_cao", "Hai biến thể màu của quần jogger", "AI_A"),
    12492: row("set", "good", "yes", "đỏ;tím;trắng", "tay_dài", "cổ_tròn", "hoa", "suông", "other", "ống_rộng", "not_visible", "Hai biến thể màu của bộ áo dài truyền thống và quần trắng", "AI_A"),
    8558: row("skirt", "good", "yes", "đen;đỏ", NA, NA, "trơn", "ôm", "midi", NA, "cạp_cao", "Hai biến thể màu của chân váy đuôi cá", "AI_A"),
    1858: row("jacket", "good", "yes", "nâu", "tay_dài", "other", "trơn", "rộng", "other", NA, NA, "Áo phao được phối với áo croptop và quần denim", "AI_A"),
    5633: row("bag", "good", "yes", "unknown", NA, NA, "trơn", NA, NA, NA, NA, "Nhiều biến thể màu của túi xách nhỏ", "AI_A"),
    5352: row("bag", "good", "no", "đen", NA, NA, "trơn", NA, NA, NA, NA, "Túi da nhỏ có dây đeo tháo rời", "AI_A"),
    6969: row("dress", "good", "yes", "be", "tay_dài", "cổ_v", "trơn", "ôm", "maxi", NA, "other", "Váy len dài kèm khăn choàng", "AI_A"),
    6148: row("hat", "good", "yes", "unknown", NA, NA, "trơn", NA, NA, NA, NA, "Nhiều biến thể màu của mũ bucket", "AI_A"),
    1774: row("jacket", "good", "no", "xanh_navy", "tay_dài", "cổ_bẻ", "trơn", "suông", "other", NA, NA, "Áo khoác denim màu xanh đậm", "AI_A"),
    9690: row("top", "good", "yes", "trắng;xanh_dương;hồng", "tay_ngắn", "cổ_bẻ", "kẻ_sọc", "rộng", "other", NA, NA, "Hai biến thể màu của áo sơ mi kẻ sọc", "AI_A"),
    4017: row("set", "good", "yes", "trắng;hồng;đen", "tay_ngắn", "cổ_v", "in_hình", "suông", "other", NA, "unknown", "Bộ áo thể thao và quần ngắn", "AI_A"),
    8199: row("skirt", "good", "yes", "xám", NA, NA, "trơn", "suông", "mini", NA, "cạp_cao", "Chân váy ánh kim được phối với áo và túi", "AI_A"),
    10304: row("top", "usable", "yes", "hồng", "tay_ngắn", "cổ_vuông", "hoa", "ôm", "croptop", NA, NA, "Áo croptop có phần ren hoa và nơ trước ngực", "AI_A"),
    466: row("top", "usable", "yes", "xám", "tay_dài", "cổ_bẻ", "trơn", "suông", "other", NA, NA, "Áo sơ mi được phối với quần trắng", "AI_A"),
    5554: row("bag", "good", "yes", "unknown", NA, NA, "trơn", NA, NA, NA, NA, "Nhiều biến thể màu của túi đeo hông", "AI_A"),
    6799: row("dress", "poor", "yes", "hồng", "not_visible", "not_visible", "hoa", "not_visible", "maxi", NA, "other", "Ảnh chỉ hiển thị phần dưới của váy; thân trên không quan sát được", "AI_A"),
    5863: row("hat", "good", "yes", "unknown", NA, NA, "trơn", NA, NA, NA, NA, "Nhiều biến thể màu của mũ len", "AI_A"),
    2629: row("jacket", "good", "yes", "xanh_dương;trắng", "tay_dài", "other", "in_hình", "rộng", "other", NA, NA, "Áo khoác varsity được trình bày ở nhiều góc nhìn", "AI_A"),
    1226: row("pants", "good", "yes", "đen", NA, NA, "trơn", "rộng", NA, NA, "cạp_cao", "Quần cargo ngắn trẻ em; loại ống quần không áp dụng", "AI_A"),
    4250: row("set", "good", "yes", "unknown", "không_tay", "cổ_tròn", "trơn", "ôm", "croptop", "ống_bó", "cạp_cao", "Nhiều biến thể màu của bộ đồ tập", "AI_A"),
    8018: row("skirt", "good", "no", "xanh_lá", NA, NA, "trơn", "suông", "midi", NA, "cạp_cao", "Chân váy nhiều tầng bèo", "AI_A"),
    9230: row("top", "good", "yes", "xanh_lá", "tay_ngắn", "cổ_tròn", "trơn", "suông", "other", NA, NA, "Cùng một áo thun ở ảnh sản phẩm và trên người mẫu", "AI_A"),
    12472: row("top", "usable", "yes", "nâu;đỏ", "tay_dài", "cổ_bẻ", "trơn", "suông", "other", NA, NA, "Nhiều biến thể màu của áo sơ mi satin", "AI_A"),
    5135: row("bag", "good", "yes", "đen", NA, NA, "trơn", NA, NA, NA, NA, "Túi đeo chéo được chụp trên người mẫu", "AI_A"),
    7805: row("dress", "good", "yes", "đen;trắng", "không_tay", "other", "hoa", "ôm", "maxi", NA, "other", "Váy dài cổ đứng có họa tiết hoa", "AI_A"),
    5904: row("hat", "usable", "yes", "đen", NA, NA, "in_hình", NA, NA, NA, NA, "Mũ lưỡi trai có chữ được chụp cùng nhiều trang phục và túi", "AI_A"),
    1790: row("jacket", "usable", "yes", "nâu", "tay_dài", "other", "trơn", "rộng", "other", NA, NA, "Áo khoác dáng rộng được phối với nhiều lớp trang phục", "AI_A"),
    12045: row("pants", "good", "yes", "tím", NA, NA, "trơn", "suông", NA, "ống_loe", "cạp_cao", "Quần dài được chụp trên người mẫu", "AI_A"),
}


def validate_annotations():
    for product_id, values in ANNOTATIONS.items():
        for column, allowed in ALLOWED.items():
            value = values[column]
            if value in (None, ""):
                raise ValueError(f"Blank required label: product_id={product_id}, column={column}")
            invalid = set(str(value).split(";")) - allowed
            if invalid:
                raise ValueError(
                    f"Invalid value(s) {sorted(invalid)}: product_id={product_id}, column={column}"
                )


def main():
    path = Path("outputs/prepared/blind_506_before_AI_A.xlsx")
    backup = path.with_name("blind_506_before_AI_A_original.xlsx")
    validate_annotations()
    workbook = load_workbook(path)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    expected = ["product_id", "split", "image_path", *COLUMNS]
    if headers != expected:
        raise ValueError(f"Unexpected workbook columns: {headers}")
    product_rows = {int(sheet.cell(r, 1).value): r for r in range(2, sheet.max_row + 1)}
    if len(product_rows) != 50 or set(product_rows) != set(ANNOTATIONS):
        missing = sorted(set(product_rows) - set(ANNOTATIONS))
        extra = sorted(set(ANNOTATIONS) - set(product_rows))
        raise ValueError(f"Expected the reviewed 50 rows; missing={missing}, extra={extra}")
    if not backup.exists():
        shutil.copy2(path, backup)
    column_index = {name: index + 1 for index, name in enumerate(headers)}
    for product_id, values in ANNOTATIONS.items():
        excel_row = product_rows[product_id]
        for column, value in values.items():
            sheet.cell(excel_row, column_index[column]).value = value
    temporary = path.with_name("blind_506_before_AI_A.__tmp__.xlsx")
    workbook.save(temporary)
    output = path
    try:
        temporary.replace(path)
    except PermissionError:
        output = path.with_name("blind_506_after_AI_A.xlsx")
        temporary.replace(output)
        print(f"Source workbook is locked; wrote the completed workbook to {output}")
    print(f"Filled {len(ANNOTATIONS)} rows in {output}")
    print(f"Backup: {backup}")


if __name__ == "__main__":
    main()
