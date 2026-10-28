"""Fill outputs/prepared/blind_50.xlsx from an image-only AI_A review."""
from pathlib import Path
import shutil

from openpyxl import load_workbook


COLUMNS = [
    "category_gold", "image_quality", "multiple_products", "color_gold",
    "sleeve_gold", "neckline_gold", "pattern_gold", "fit_gold", "length_gold",
    "pant_cut_gold", "waist_gold", "notes", "annotator",
]


def row(*values):
    if len(values) != len(COLUMNS):
        raise ValueError(f"Expected {len(COLUMNS)} values, got {len(values)}")
    return dict(zip(COLUMNS, values))


NA = "not_applicable"

ANNOTATIONS = {
    6658: row("dress", "good", "no", "xanh_navy", "tay_ngắn", "cổ_tròn", "trơn", "ôm", "midi", NA, "other", "Phần thân trên có chi tiết ren", "AI_A"),
    6251: row("hat", "good", "no", "be", NA, NA, "trơn", NA, NA, NA, NA, "", "AI_A"),
    1946: row("jacket", "usable", "no", "xám", "tay_dài", "other", "trơn", "rộng", "other", NA, NA, "Áo khoác có mũ; chữ quảng cáo không che sản phẩm chính", "AI_A"),
    12083: row("set", "usable", "yes", "vàng;trắng", "tay_dài", "cổ_tròn", "trơn", "rộng", "other", "ống_rộng", "unknown", "Một bộ trang phục được lặp lại ở nhiều khung hình", "AI_A"),
    3973: row("set", "good", "yes", "đen;trắng", "tay_ngắn", "cổ_tròn", "trơn", "rộng", "other", NA, "unknown", "Bộ áo thun và quần ngắn", "AI_A"),
    8287: row("top", "usable", "yes", "xanh_lá;đen", "không_tay", "hai_dây", "trơn", "suông", "other", NA, NA, "Gán thuộc tính cho áo hai dây; ảnh có phối thêm chân váy", "AI_A"),
    329: row("top", "good", "no", "trắng", "tay_ngắn", "cổ_tròn", "in_hình", "rộng", "other", NA, NA, "", "AI_A"),
    5014: row("bag", "good", "no", "đỏ", NA, NA, "other", NA, NA, NA, NA, "Bề mặt túi chần ô ngoài vocabulary pattern", "AI_A"),
    5569: row("bag", "good", "no", "hồng", NA, NA, "other", NA, NA, NA, NA, "Bề mặt túi chần quả trám ngoài vocabulary pattern", "AI_A"),
    7037: row("dress", "usable", "no", "hồng", "tay_dài", "trễ_vai", "trơn", "suông", "mini", NA, "other", "", "AI_A"),
    6274: row("hat", "usable", "yes", "unknown", NA, NA, "trơn", NA, NA, NA, NA, "Nhiều biến thể màu của mũ", "AI_A"),
    4818: row("jacket", "good", "no", "trắng", "tay_dài", "cổ_bẻ", "trơn", "ôm", "other", NA, NA, "", "AI_A"),
    933: row("pants", "good", "no", "xanh_dương", NA, NA, "trơn", "ôm", NA, "ống_bó", "cạp_cao", "Quần denim có hiệu ứng bạc màu nhẹ", "AI_A"),
    9554: row("set", "usable", "yes", "kem", "tay_dài", "cổ_v", "hoa", "suông", "croptop", NA, NA, "Bộ áo hai dây kèm cardigan", "AI_A"),
    8050: row("skirt", "usable", "yes", "xám", NA, NA, "trơn", "suông", "mini", NA, "cạp_cao", "Cùng một chân váy được trình bày trên người mẫu và ảnh sản phẩm", "AI_A"),
    710: row("top", "good", "yes", "đỏ", "tay_dài", "cổ_bẻ", "trơn", "ôm", "other", NA, NA, "Ảnh phối thêm chân váy trắng", "AI_A"),
    4152: row("top", "usable", "yes", "xanh_dương", "tay_ngắn", "cổ_bẻ", "trơn", "ôm", "other", NA, NA, "Cùng một áo ở hai góc nhìn", "AI_A"),
    5430: row("bag", "good", "no", "be", NA, NA, "caro", NA, NA, NA, NA, "Có móc trang trí gắn trên túi", "AI_A"),
    6702: row("top", "usable", "yes", "vàng", "tay_ngắn", "cổ_tròn", "in_hình", "rộng", "other", NA, NA, "Cùng một áo ở mặt trước và sau", "AI_A"),
    11272: row("jacket", "good", "no", "đỏ", "không_tay", "other", "kẻ_sọc", "rộng", "other", NA, NA, "Áo choàng poncho có mũ và tua rua", "AI_A"),
    11471: row("jacket", "usable", "yes", "unknown", "tay_dài", "other", "trơn", "rộng", "other", NA, NA, "Nhiều biến thể màu của áo hoodie", "AI_A"),
    12860: row("set", "good", "yes", "hồng;trắng", "tay_dài", "cổ_lọ", "hoa", "ôm", "maxi", "ống_rộng", "not_visible", "Bộ áo dài kèm quần", "AI_A"),
    4448: row("set", "good", "yes", "đen;xanh_dương", "tay_ngắn", "cổ_tròn", "in_hình", "suông", "other", NA, "unknown", "Bộ áo thể thao và quần ngắn", "AI_A"),
    8316: row("set", "usable", "yes", "trắng;xanh_dương", "tay_ngắn", "cổ_tròn", "trơn", "suông", "mini", NA, "cạp_cao", "Ảnh phối áo thun với chân váy denim", "AI_A"),
    750: row("top", "good", "no", "trắng", "tay_ngắn", "cổ_tròn", "in_hình", "rộng", "other", NA, NA, "", "AI_A"),
    4400: row("set", "usable", "yes", "unknown", "tay_ngắn", "cổ_tròn", "in_hình", "rộng", "other", NA, "unknown", "Hai biến thể màu của bộ áo và quần ngắn", "AI_A"),
    3127: row("dress", "usable", "yes", "unknown", "tay_dài", "cổ_v", "trơn", "ôm", "mini", NA, "other", "Nhiều biến thể màu của cùng kiểu váy", "AI_A"),
    5932: row("fabric", "usable", "yes", "unknown", NA, NA, "trơn", NA, NA, NA, NA, "Nhiều mảnh vải hoặc sản phẩm gấp với nhiều màu", "AI_A"),
    1861: row("jacket", "usable", "yes", "trắng", "tay_dài", "other", "trơn", "rộng", "other", NA, NA, "Cùng kiểu áo khoác trên hai người mẫu", "AI_A"),
    11742: row("pants", "good", "no", "xám", NA, NA, "trơn", "rộng", NA, "ống_rộng", "cạp_cao", "Nếp ly dọc là cấu trúc vải không phải họa tiết kẻ", "AI_A"),
    4183: row("set", "good", "no", "xanh_lá", "tay_ngắn", "cổ_tròn", "trơn", "ôm", NA, NA, "cạp_cao", "Trang phục liền quần ngắn ngoài taxonomy", "AI_A"),
    8434: row("set", "good", "yes", "trắng", "không_tay", "trễ_vai", "trơn", "ôm", "maxi", NA, "cạp_cao", "Bộ áo croptop và chân váy dài", "AI_A"),
    9633: row("top", "good", "yes", "trắng", "tay_ngắn", "cổ_lọ", "trơn", "ôm", "croptop", NA, NA, "Ảnh có minh họa thêm các biến thể màu", "AI_A"),
    5705: row("bag", "good", "no", "nâu", NA, NA, "hoa", NA, NA, NA, NA, "Hoa đính nổi trên túi đan", "AI_A"),
    5503: row("bag", "good", "no", "nâu", NA, NA, "trơn", NA, NA, NA, NA, "", "AI_A"),
    6808: row("dress", "usable", "no", "kem", "không_tay", "hai_dây", "trơn", "suông", "maxi", NA, "other", "", "AI_A"),
    6000: row("hat", "usable", "yes", "unknown", NA, NA, "trơn", NA, NA, NA, NA, "Nhiều biến thể màu của mũ lưỡi trai", "AI_A"),
    4736: row("jacket", "good", "no", "xám", "tay_dài", "cổ_bẻ", "caro", "ôm", "other", NA, NA, "", "AI_A"),
    12820: row("dress", "usable", "yes", "xanh_lá", "not_visible", "not_visible", "trơn", "not_visible", "not_visible", NA, "not_visible", "Sản phẩm áo dài được gấp nên không quan sát được cấu trúc", "AI_A"),
    12542: row("set", "usable", "yes", "đỏ;tím;trắng", "tay_lỡ", "cổ_tròn", "hoa", "ôm", "maxi", "ống_rộng", "not_visible", "Hai biến thể áo dài kèm quần trắng", "AI_A"),
    8590: row("skirt", "good", "no", "trắng;đen", NA, NA, "hoa", "suông", "midi", NA, "cạp_cao", "", "AI_A"),
    9949: row("top", "usable", "yes", "xanh_lá;trắng", "tay_dài", "cổ_bẻ", "kẻ_sọc", "rộng", "other", NA, NA, "Ảnh có nhiều biến thể màu nhỏ ở bên phải", "AI_A"),
    4254: row("set", "good", "yes", "xanh_dương", "tay_ngắn", "cổ_tròn", "trơn", "ôm", "croptop", "ống_bó", "cạp_cao", "Bộ áo croptop và quần legging", "AI_A"),
    5159: row("bag", "good", "no", "xanh_lá", NA, NA, "trơn", NA, NA, NA, NA, "", "AI_A"),
    6748: row("dress", "good", "no", "xanh_lá", "không_tay", "cổ_tròn", "trơn", "suông", "mini", NA, "other", "", "AI_A"),
    5768: row("hat", "good", "yes", "xám;đen;đỏ", NA, NA, "in_hình", NA, NA, NA, NA, "Hai biến thể màu có hình Người Nhện", "AI_A"),
    2606: row("jacket", "usable", "yes", "đen", "tay_dài", "cổ_lọ", "kẻ_sọc", "rộng", "other", NA, NA, "Cùng một áo khoác ở nhiều khung hình", "AI_A"),
    13212: row("set", "good", "yes", "trắng;đen", "tay_ngắn", "cổ_tròn", "in_hình", "rộng", "other", NA, "unknown", "Bộ áo thun và quần ngắn", "AI_A"),
    13090: row("set", "usable", "yes", "đen", "tay_dài", "cổ_lọ", "trơn", "suông", "other", "ống_suông", "not_visible", "Bộ trang phục truyền thống gồm áo và quần", "AI_A"),
    9006: row("set", "usable", "yes", "xám;đen", "tay_dài", "cổ_tròn", "caro", "suông", "mini", NA, "cạp_cao", "Ảnh phối áo len với chân váy caro", "AI_A"),
}


def main():
    path = Path("outputs/prepared/blind_50.xlsx")
    backup = path.with_name("blind_50_before_AI_A.xlsx")
    workbook = load_workbook(path)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    expected = ["product_id", "split", "image_path", *COLUMNS]
    if headers != expected:
        raise ValueError(f"Unexpected workbook columns: {headers}")
    product_rows = {int(sheet.cell(r, 1).value): r for r in range(2, sheet.max_row + 1)}
    if len(product_rows) != 50 or set(product_rows) != set(ANNOTATIONS):
        missing = sorted(set(product_rows) - set(ANNOTATIONS))
        extra = sorted(set(ANNOTATIONS) - set(product_rows))
        raise ValueError(f"Expected the reviewed 50 rows; missing={missing}, extra={extra}")
    if not backup.exists():
        shutil.copy2(path, backup)
    column_index = {name: index + 1 for index, name in enumerate(headers)}
    for product_id, values in ANNOTATIONS.items():
        excel_row = product_rows[product_id]
        for column, value in values.items():
            sheet.cell(excel_row, column_index[column]).value = value
    temporary = path.with_name("blind_50.__tmp__.xlsx")
    workbook.save(temporary)
    temporary.replace(path)
    print(f"Filled {len(ANNOTATIONS)} rows in {path}")
    print(f"Backup: {backup}")


if __name__ == "__main__":
    main()
